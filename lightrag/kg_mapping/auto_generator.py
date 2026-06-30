"""Automatic KG mapping generation from relational database metadata."""

from __future__ import annotations

import json
import re
import asyncio
import csv
import hashlib
import io
import shutil
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from fnmatch import fnmatchcase
from pathlib import Path
from typing import Any

import yaml


_IDENTIFIER_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_TEXT_TYPES = {"text", "varchar", "character varying", "char", "character"}
_EVENT_HINTS = ("record", "bid", "contract", "order", "payment", "apply", "approval")
_EVENT_ATTR_HINTS = ("amount", "date", "time", "status", "rank", "price", "total")


@dataclass(frozen=True)
class ColumnInfo:
    table_name: str
    column_name: str
    data_type: str
    is_nullable: bool = True
    comment: str | None = None


@dataclass(frozen=True)
class TableInfo:
    table_name: str
    columns: list[ColumnInfo]
    primary_key: str | None = None
    comment: str | None = None
    row_count: int | None = None


@dataclass(frozen=True)
class RelationshipCandidate:
    source_table: str
    source_column: str
    target_table: str
    target_column: str
    source: str
    score: float
    evidence: dict[str, Any]
    decision: str


def infer_relationships(
    tables: list[TableInfo],
    *,
    explicit_relationships: list[RelationshipCandidate] | None = None,
    coverage: dict[tuple[str, str, str, str], float] | None = None,
    auto_approve_threshold: float = 0.85,
    review_threshold: float = 0.65,
) -> list[RelationshipCandidate]:
    """Infer table relationships from explicit metadata, names, and data coverage."""

    table_by_name = {table.table_name: table for table in tables}
    by_key: dict[tuple[str, str, str, str], RelationshipCandidate] = {}
    for relationship in explicit_relationships or []:
        key = (
            relationship.source_table,
            relationship.source_column,
            relationship.target_table,
            relationship.target_column,
        )
        by_key[key] = relationship

    coverage = coverage or {}
    for source_table in tables:
        for column in source_table.columns:
            if not _looks_like_reference(column.column_name):
                continue
            for target_table in tables:
                if target_table.table_name == source_table.table_name:
                    continue
                target_column_name = _matching_target_column(
                    column.column_name,
                    target_table,
                )
                if not target_column_name:
                    continue
                key = (
                    source_table.table_name,
                    column.column_name,
                    target_table.table_name,
                    target_column_name,
                )
                if key in by_key:
                    continue

                target_column = _column_by_name(target_table, target_column_name)
                score, evidence = _score_inferred_relationship(
                    column,
                    target_table,
                    target_column,
                    coverage.get(key),
                )
                if score < review_threshold:
                    continue
                decision = (
                    "auto_approved"
                    if score >= auto_approve_threshold
                    else "need_review"
                )
                by_key[key] = RelationshipCandidate(
                    source_table=source_table.table_name,
                    source_column=column.column_name,
                    target_table=target_table.table_name,
                    target_column=target_column_name,
                    source="inferred_by_name_and_data",
                    score=round(min(score, 1.0), 4),
                    evidence=evidence,
                    decision=decision,
                )

    return sorted(
        by_key.values(),
        key=lambda item: (
            item.source_table,
            item.source_column,
            item.target_table,
            item.target_column,
        ),
    )


def generate_mapping_from_schema(
    *,
    database_name: str,
    tables: list[TableInfo],
    relationships: list[RelationshipCandidate],
) -> dict[str, Any]:
    """Generate a conservative `kg_mappings` YAML object from schema metadata."""

    table_by_name = {table.table_name: table for table in tables}
    approved_relationships = [
        relationship
        for relationship in relationships
        if relationship.decision == "auto_approved"
    ]
    relationships_by_source: dict[str, list[RelationshipCandidate]] = {}
    for relationship in approved_relationships:
        relationships_by_source.setdefault(relationship.source_table, []).append(
            relationship
        )

    entity_types: dict[str, dict[str, str]] = {}
    sources: dict[str, dict[str, Any]] = {}
    entities: list[dict[str, Any]] = []
    mapping_relationships: list[dict[str, Any]] = []

    for table in tables:
        if not table.primary_key:
            continue
        source_relationships = relationships_by_source.get(table.table_name, [])
        is_event = _is_event_table(table, source_relationships)
        sources[table.table_name] = _source_config_for_table(
            table,
            source_relationships,
            table_by_name,
        )

        entity_type = _entity_type_for_table(table, is_event=is_event)
        entity_types.setdefault(
            entity_type,
            {
                "label": _label_for_entity_type(entity_type),
                "id_prefix": entity_type,
            },
        )
        entity_config = _entity_config_for_table(
            table,
            entity_type,
            source_relationships,
            table_by_name,
            is_event=is_event,
        )
        entities.append(entity_config)

    for relationship in approved_relationships:
        source_table = table_by_name.get(relationship.source_table)
        target_table = table_by_name.get(relationship.target_table)
        if not source_table or not target_table:
            continue
        if not source_table.primary_key or not target_table.primary_key:
            continue

        source_entity_type = _entity_type_for_table(
            source_table,
            is_event=_is_event_table(
                source_table,
                relationships_by_source.get(source_table.table_name, []),
            ),
        )
        target_entity_type = _entity_type_for_table(target_table, is_event=False)
        relation_type = _relation_type(relationship.source_table, relationship.target_table)
        display_field = _display_field(target_table)
        display_placeholder = (
            _join_output_name(relationship, target_table, display_field)
            if display_field
            else relationship.source_column
        )
        mapping_relationships.append(
            {
                "source": relationship.source_table,
                "relation_type": relation_type,
                "src": {
                    "entity_type": source_entity_type,
                    "id_field": source_table.primary_key,
                },
                "tgt": {
                    "entity_type": target_entity_type,
                    "id_field": relationship.source_column,
                },
                "description_template": (
                    "{"
                    + source_table.primary_key
                    + "} "
                    + relation_type
                    + " {"
                    + display_placeholder
                    + "}."
                ),
                "x_inference": _relationship_as_metadata(relationship),
            }
        )

    return {
        "schema_version": "auto_kg_v1",
        "database_name": database_name,
        "sources": sources,
        "entity_types": entity_types,
        "entities": entities,
        "relationships": mapping_relationships,
    }


def quote_ident(identifier: str) -> str:
    """Quote a simple SQL identifier."""

    if not _IDENTIFIER_RE.match(identifier):
        raise ValueError(f"Invalid SQL identifier: {identifier}")
    return f'"{identifier}"'


def coverage_query(
    schema: str,
    source_table: str,
    source_column: str,
    target_table: str,
    target_column: str,
    sample_limit: int,
) -> str:
    """Build a coverage query for validating inferred relationships."""

    if sample_limit < 1:
        raise ValueError("sample_limit must be >= 1")
    source_ref = f"{quote_ident(schema)}.{quote_ident(source_table)}"
    target_ref = f"{quote_ident(schema)}.{quote_ident(target_table)}"
    source_col = quote_ident(source_column)
    target_col = quote_ident(target_column)
    return (
        "WITH sampled AS ("
        f"SELECT {source_col} AS value FROM {source_ref} "
        f"WHERE {source_col} IS NOT NULL LIMIT {sample_limit}"
        ") "
        "SELECT count(*) AS total, "
        f"count({target_ref}.{target_col}) AS matched "
        f"FROM sampled LEFT JOIN {target_ref} "
        f"ON sampled.value = {target_ref}.{target_col}"
    )


def new_generation_id(now: datetime | None = None) -> str:
    timestamp = (now or datetime.now(timezone.utc)).strftime("%Y%m%d%H%M%S%f")
    return f"gen_{timestamp}"


def save_generation_record(record_dir: str | Path, record: dict[str, Any]) -> Path:
    generation_id = str(record["generation_id"])
    path = Path(record_dir) / f"{generation_id}.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as file:
        json.dump(record, file, ensure_ascii=False, indent=2)
        file.write("\n")
    return path


def load_generation_record(record_dir: str | Path, generation_id: str) -> dict[str, Any]:
    path = Path(record_dir) / f"{generation_id}.json"
    with path.open("r", encoding="utf-8") as file:
        data = json.load(file)
    if not isinstance(data, dict):
        raise ValueError(f"Generation record must be a JSON object: {path}")
    return data


def write_mapping_yaml(
    mapping_dir: str | Path,
    database_name: str,
    generation_id: str,
    mapping: dict[str, Any],
) -> Path:
    path = Path(mapping_dir) / f"{database_name}.{generation_id}.yaml"
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as file:
        yaml.safe_dump(mapping, file, allow_unicode=True, sort_keys=False)
    return path


def write_current_mapping(
    mapping_dir: str | Path,
    database_name: str,
    mapping_path: str | Path,
) -> Path:
    path = Path(mapping_dir) / f"{database_name}.current.yaml"
    path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(mapping_path, path)
    return path


def parse_relationship_metadata_file(
    filename: str,
    content: bytes,
) -> list[RelationshipCandidate]:
    """Parse customer-supplied ER relationship metadata."""

    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        data = json.loads(content.decode("utf-8-sig"))
        return _relationships_from_records(_relationship_records(data), filename)
    if suffix in {".yaml", ".yml"}:
        data = yaml.safe_load(content.decode("utf-8-sig"))
        return _relationships_from_records(_relationship_records(data), filename)
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        records = list(csv.DictReader(io.StringIO(text)))
        return _relationships_from_records(records, filename)
    if suffix in {".sql", ".ddl"}:
        return _relationships_from_ddl(content.decode("utf-8-sig"), filename)
    if suffix == ".xlsx":
        return _relationships_from_xlsx(content, filename)
    raise ValueError(f"Unsupported relationship metadata file type: {filename}")


def metadata_file_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def merge_mapping_configs(
    base_mapping: dict[str, Any] | None,
    generated_mapping: dict[str, Any],
) -> dict[str, Any]:
    """Merge generated mapping sections into an existing mapping by source."""

    if not base_mapping:
        return generated_mapping

    merged = {
        "schema_version": generated_mapping.get(
            "schema_version",
            base_mapping.get("schema_version"),
        ),
        "database_name": generated_mapping.get(
            "database_name",
            base_mapping.get("database_name", "default"),
        ),
        "sources": {
            **base_mapping.get("sources", {}),
            **generated_mapping.get("sources", {}),
        },
        "entity_types": {
            **base_mapping.get("entity_types", {}),
            **generated_mapping.get("entity_types", {}),
        },
        "entities": _merge_list_by_source(
            base_mapping.get("entities", []),
            generated_mapping.get("entities", []),
        ),
        "relationships": _merge_relationships(
            base_mapping.get("relationships", []),
            generated_mapping.get("relationships", []),
        ),
    }
    for key, value in generated_mapping.items():
        if key not in merged:
            merged[key] = value
    return merged


def apply_llm_mapping_enhancement(
    mapping: dict[str, Any],
    enhancement: dict[str, Any],
) -> tuple[dict[str, Any], dict[str, Any]]:
    """Apply LLM readability suggestions without changing structural mapping fields."""

    enhanced = json.loads(json.dumps(mapping, ensure_ascii=False))
    trace = {
        "applied": {
            "entity_labels": 0,
            "entities": 0,
            "relationships": 0,
        },
        "ignored": [],
        "raw": enhancement,
    }

    entity_types = enhanced.get("entity_types", {})
    for suggestion in _normalize_entity_label_suggestions(
        enhancement.get("entity_labels")
    ):
        if not isinstance(suggestion, dict):
            continue
        entity_type = str(suggestion.get("entity_type") or "")
        label = str(suggestion.get("label") or "").strip()
        if entity_type in entity_types and label:
            entity_types[entity_type]["label"] = label
            trace["applied"]["entity_labels"] += 1

    entities_by_source = {
        str(entity.get("source")): entity
        for entity in enhanced.get("entities", [])
        if isinstance(entity, dict)
    }
    for suggestion in enhancement.get("entities", []) or []:
        if not isinstance(suggestion, dict):
            continue
        source = str(suggestion.get("source") or "")
        entity = entities_by_source.get(source)
        if not entity:
            trace["ignored"].append({"type": "entity", "source": source})
            continue
        changed = False
        for field in ("entity_name_template", "description_template"):
            value = suggestion.get(field)
            if isinstance(value, str) and value.strip():
                entity[field] = value.strip()
                changed = True
        if changed:
            entity["x_llm_enhanced"] = True
            trace["applied"]["entities"] += 1

    relationships = [
        relationship
        for relationship in enhanced.get("relationships", [])
        if isinstance(relationship, dict)
    ]
    relationship_suggestions = [
        suggestion
        for suggestion in (enhancement.get("relationships", []) or [])
        if isinstance(suggestion, dict)
    ]
    source_offsets: dict[str, int] = {}
    for suggestion in relationship_suggestions:
        if not isinstance(suggestion, dict):
            continue
        relationship = _find_relationship_for_llm_suggestion(
            relationships,
            suggestion,
            source_offsets,
        )
        if not relationship:
            trace["ignored"].append(
                {
                    "type": "relationship",
                    "source": suggestion.get("source"),
                    "old_relation_type": suggestion.get("old_relation_type"),
                }
            )
            continue
        changed = False
        relation_type = suggestion.get("relation_type")
        if isinstance(relation_type, str) and _IDENTIFIER_RE.match(relation_type):
            relationship["relation_type"] = relation_type.upper()
            changed = True
        description = suggestion.get("description_template")
        if isinstance(description, str) and description.strip():
            relationship["description_template"] = description.strip()
            changed = True
        if changed:
            relationship["x_llm_enhanced"] = True
            trace["applied"]["relationships"] += 1

    return enhanced, trace


def _normalize_entity_label_suggestions(value: Any) -> list[dict[str, Any]]:
    if isinstance(value, dict):
        return [
            {"entity_type": entity_type, "label": label}
            for entity_type, label in value.items()
        ]
    if isinstance(value, list):
        return value
    return []


def introspect_postgres_schema(
    connection_url: str,
    *,
    schema: str,
    sample_limit: int = 500,
) -> tuple[list[TableInfo], list[RelationshipCandidate], dict[tuple[str, str, str, str], float]]:
    """Read PostgreSQL schema metadata and relationship coverage."""

    try:
        return _introspect_postgres_schema_with_psycopg(
            connection_url,
            schema=schema,
            sample_limit=sample_limit,
        )
    except RuntimeError as psycopg_error:
        try:
            return asyncio.run(
                _introspect_postgres_schema_with_asyncpg(
                    connection_url,
                    schema=schema,
                    sample_limit=sample_limit,
                )
            )
        except ImportError as asyncpg_error:
            raise psycopg_error from asyncpg_error


def _introspect_postgres_schema_with_psycopg(
    connection_url: str,
    *,
    schema: str,
    sample_limit: int,
) -> tuple[list[TableInfo], list[RelationshipCandidate], dict[tuple[str, str, str, str], float]]:
    try:
        import psycopg  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "PostgreSQL KG mapping generation requires 'psycopg' or 'asyncpg'."
        ) from exc

    with psycopg.connect(connection_url) as connection:
        tables = _load_postgres_tables(connection, schema)
        explicit = _load_postgres_foreign_keys(connection, schema)
        inferred_seed = infer_relationships(
            tables,
            explicit_relationships=explicit,
            review_threshold=0.0,
        )
        coverage = _load_relationship_coverage(
            connection,
            schema,
            inferred_seed,
            sample_limit,
        )
    return tables, explicit, coverage


async def _introspect_postgres_schema_with_asyncpg(
    connection_url: str,
    *,
    schema: str,
    sample_limit: int,
) -> tuple[list[TableInfo], list[RelationshipCandidate], dict[tuple[str, str, str, str], float]]:
    import asyncpg  # type: ignore

    connection = await asyncpg.connect(connection_url)
    try:
        tables = await _load_postgres_tables_asyncpg(connection, schema)
        explicit = await _load_postgres_foreign_keys_asyncpg(connection, schema)
        inferred_seed = infer_relationships(
            tables,
            explicit_relationships=explicit,
            review_threshold=0.0,
        )
        coverage = await _load_relationship_coverage_asyncpg(
            connection,
            schema,
            inferred_seed,
            sample_limit,
        )
        return tables, explicit, coverage
    finally:
        await connection.close()


def filter_schema_for_generation(
    tables: list[TableInfo],
    relationships: list[RelationshipCandidate],
    coverage: dict[tuple[str, str, str, str], float],
    *,
    excluded_tables: set[str] | None = None,
    excluded_table_patterns: list[str] | None = None,
) -> tuple[
    list[TableInfo],
    list[RelationshipCandidate],
    dict[tuple[str, str, str, str], float],
    list[str],
]:
    """Remove tables that should not be converted into generated KG mappings."""

    excluded_tables = excluded_tables or set()
    excluded_table_patterns = excluded_table_patterns or []
    excluded = {
        table.table_name
        for table in tables
        if table.table_name in excluded_tables
        or any(
            fnmatchcase(table.table_name, pattern)
            for pattern in excluded_table_patterns
        )
    }
    if not excluded:
        return tables, relationships, coverage, []

    filtered_tables = [
        table
        for table in tables
        if table.table_name not in excluded
    ]
    filtered_relationships = [
        relationship
        for relationship in relationships
        if relationship.source_table not in excluded
        and relationship.target_table not in excluded
    ]
    filtered_coverage = {
        key: value
        for key, value in coverage.items()
        if key[0] not in excluded and key[2] not in excluded
    }
    return filtered_tables, filtered_relationships, filtered_coverage, sorted(excluded)


def generation_summary(
    tables: list[TableInfo],
    mapping: dict[str, Any],
    relationships: list[RelationshipCandidate],
) -> dict[str, int]:
    auto_approved = sum(1 for item in relationships if item.decision == "auto_approved")
    need_review = sum(1 for item in relationships if item.decision == "need_review")
    blocked = sum(1 for item in relationships if item.decision == "blocked")
    return {
        "tables": len(tables),
        "entities": len(mapping.get("entities", [])),
        "relationships": len(mapping.get("relationships", [])),
        "auto_approved": auto_approved + len(mapping.get("entities", [])),
        "need_review": need_review,
        "blocked": blocked,
    }


def serialize_relationships(
    relationships: list[RelationshipCandidate],
) -> list[dict[str, Any]]:
    return [asdict(relationship) for relationship in relationships]


def _relationship_records(data: Any) -> list[dict[str, Any]]:
    if isinstance(data, dict):
        data = data.get("relationships") or data.get("relations") or []
    if not isinstance(data, list):
        raise ValueError("Relationship metadata must be a list or contain relationships")
    records = []
    for item in data:
        if not isinstance(item, dict):
            raise ValueError("Relationship metadata rows must be objects")
        records.append(item)
    return records


def _relationships_from_records(
    records: list[dict[str, Any]],
    filename: str,
) -> list[RelationshipCandidate]:
    relationships = []
    for index, record in enumerate(records, start=1):
        source_table = _record_value(
            record,
            "source_table",
            "from_table",
            "child_table",
            "table",
        )
        source_column = _record_value(
            record,
            "source_column",
            "from_column",
            "child_column",
            "column",
        )
        target_table = _record_value(
            record,
            "target_table",
            "to_table",
            "parent_table",
            "ref_table",
            "referenced_table",
        )
        target_column = _record_value(
            record,
            "target_column",
            "to_column",
            "parent_column",
            "ref_column",
            "referenced_column",
        )
        if not all([source_table, source_column, target_table, target_column]):
            raise ValueError(
                f"Relationship metadata row {index} in {filename} is incomplete"
            )
        relationships.append(
            RelationshipCandidate(
                source_table=source_table,
                source_column=source_column,
                target_table=target_table,
                target_column=target_column,
                source="er_declared",
                score=1.0,
                evidence={
                    "filename": filename,
                    "row": index,
                    "relationship_name": _record_value(
                        record,
                        "relationship_name",
                        "relation_name",
                        "name",
                    ),
                },
                decision="auto_approved",
            )
        )
    return relationships


def _record_value(record: dict[str, Any], *keys: str) -> str:
    normalized = {
        str(key).strip().lower(): value
        for key, value in record.items()
    }
    for key in keys:
        value = normalized.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _relationships_from_ddl(ddl: str, filename: str) -> list[RelationshipCandidate]:
    relationships = []
    create_table_pattern = re.compile(
        r"CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?(?P<table>[A-Za-z_][A-Za-z0-9_]*)\s*\((?P<body>.*?)\)\s*;",
        re.IGNORECASE | re.DOTALL,
    )
    inline_fk_pattern = re.compile(
        r"FOREIGN\s+KEY\s*\(\s*(?P<source_column>[A-Za-z_][A-Za-z0-9_]*)\s*\)\s*"
        r"REFERENCES\s+(?P<target_table>[A-Za-z_][A-Za-z0-9_]*)\s*"
        r"\(\s*(?P<target_column>[A-Za-z_][A-Za-z0-9_]*)\s*\)",
        re.IGNORECASE,
    )
    alter_fk_pattern = re.compile(
        r"ALTER\s+TABLE\s+(?P<source_table>[A-Za-z_][A-Za-z0-9_]*)\s+.*?"
        r"FOREIGN\s+KEY\s*\(\s*(?P<source_column>[A-Za-z_][A-Za-z0-9_]*)\s*\)\s*"
        r"REFERENCES\s+(?P<target_table>[A-Za-z_][A-Za-z0-9_]*)\s*"
        r"\(\s*(?P<target_column>[A-Za-z_][A-Za-z0-9_]*)\s*\)",
        re.IGNORECASE | re.DOTALL,
    )
    relationship_hint_pattern = re.compile(
        r"\b(?P<source_table>[A-Za-z_][A-Za-z0-9_]*)\."
        r"(?P<source_column>[A-Za-z_][A-Za-z0-9_]*)\s*"
        r"(?:->|=>|references)\s*"
        r"(?P<target_table>[A-Za-z_][A-Za-z0-9_]*)\."
        r"(?P<target_column>[A-Za-z_][A-Za-z0-9_]*)\b",
        re.IGNORECASE,
    )
    for create_match in create_table_pattern.finditer(ddl):
        source_table = create_match.group("table")
        for fk_match in inline_fk_pattern.finditer(create_match.group("body")):
            relationships.append(
                _ddl_relationship(
                    source_table,
                    fk_match.group("source_column"),
                    fk_match.group("target_table"),
                    fk_match.group("target_column"),
                    filename,
                )
            )
    for fk_match in alter_fk_pattern.finditer(ddl):
        relationships.append(
            _ddl_relationship(
                fk_match.group("source_table"),
                fk_match.group("source_column"),
                fk_match.group("target_table"),
                fk_match.group("target_column"),
                filename,
            )
        )
    for hint_match in relationship_hint_pattern.finditer(ddl):
        relationships.append(
            _ddl_relationship(
                hint_match.group("source_table"),
                hint_match.group("source_column"),
                hint_match.group("target_table"),
                hint_match.group("target_column"),
                filename,
            )
        )
    return _dedupe_relationships(relationships)


def _dedupe_relationships(
    relationships: list[RelationshipCandidate],
) -> list[RelationshipCandidate]:
    deduped: dict[tuple[str, str, str, str], RelationshipCandidate] = {}
    for relationship in relationships:
        key = (
            relationship.source_table,
            relationship.source_column,
            relationship.target_table,
            relationship.target_column,
        )
        deduped.setdefault(key, relationship)
    return list(deduped.values())


def _ddl_relationship(
    source_table: str,
    source_column: str,
    target_table: str,
    target_column: str,
    filename: str,
) -> RelationshipCandidate:
    return RelationshipCandidate(
        source_table=source_table,
        source_column=source_column,
        target_table=target_table,
        target_column=target_column,
        source="er_declared",
        score=1.0,
        evidence={"filename": filename, "format": "ddl"},
        decision="auto_approved",
    )


def _relationships_from_xlsx(
    content: bytes,
    filename: str,
) -> list[RelationshipCandidate]:
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:
        raise RuntimeError("Excel relationship metadata requires openpyxl") from exc

    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "").strip() for value in rows[0]]
    records = []
    for row in rows[1:]:
        if not any(value is not None and str(value).strip() for value in row):
            continue
        records.append(
            {
                header: value
                for header, value in zip(headers, row)
                if header
            }
        )
    return _relationships_from_records(records, filename)


def _merge_list_by_source(
    base_items: list[dict[str, Any]],
    generated_items: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    by_source = {str(item.get("source")): item for item in base_items}
    for item in generated_items:
        by_source[str(item.get("source"))] = item
    return list(by_source.values())


def _merge_relationships(
    base_items: list[dict[str, Any]],
    generated_items: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    by_key = {
        _relationship_config_key(item): item
        for item in base_items
    }
    for item in generated_items:
        by_key[_relationship_config_key(item)] = item
    return list(by_key.values())


def _relationship_config_key(item: dict[str, Any]) -> tuple[str, str, str, str, str]:
    return (
        str(item.get("source")),
        str(item.get("relation_type")),
        str(item.get("src", {}).get("entity_type")),
        str(item.get("src", {}).get("id_field")),
        str(item.get("tgt", {}).get("id_field")),
    )


def _find_relationship_for_llm_suggestion(
    relationships: list[dict[str, Any]],
    suggestion: dict[str, Any],
    source_offsets: dict[str, int] | None = None,
) -> dict[str, Any] | None:
    source = str(suggestion.get("source") or "")
    old_relation_type = str(suggestion.get("old_relation_type") or "")
    target_entity_type = str(suggestion.get("target_entity_type") or "")
    candidates = [
        relationship
        for relationship in relationships
        if str(relationship.get("source") or "") == source
    ]
    if old_relation_type:
        for relationship in candidates:
            if str(relationship.get("relation_type") or "") == old_relation_type:
                return relationship
    if target_entity_type:
        for relationship in candidates:
            if str(relationship.get("tgt", {}).get("entity_type") or "") == target_entity_type:
                return relationship
    if source and source_offsets is not None and candidates:
        offset = source_offsets.get(source, 0)
        if offset < len(candidates):
            source_offsets[source] = offset + 1
            return candidates[offset]
    return candidates[0] if len(candidates) == 1 else None


def _looks_like_reference(column_name: str) -> bool:
    return column_name.endswith("_id") or column_name.endswith("_code") or column_name.endswith("_no")


def _matching_target_column(column_name: str, target_table: TableInfo) -> str | None:
    candidates = []
    if target_table.primary_key:
        candidates.append(target_table.primary_key)
    candidates.extend(
        column.column_name
        for column in target_table.columns
        if column.column_name != target_table.primary_key
    )
    source_prefix = _reference_prefix(column_name)
    for candidate in candidates:
        if candidate == column_name and (
            target_table.table_name == source_prefix
            or target_table.table_name.endswith(source_prefix)
            or source_prefix.endswith(target_table.table_name)
        ):
            return candidate
        if candidate == f"{target_table.table_name}_id" and column_name == candidate:
            return candidate
        if source_prefix == target_table.table_name and candidate in {
            "id",
            f"{target_table.table_name}_id",
            column_name,
        }:
            return candidate
    return None


def _reference_prefix(column_name: str) -> str:
    for suffix in ("_id", "_code", "_no"):
        if column_name.endswith(suffix):
            return column_name[: -len(suffix)]
    return column_name


def _column_by_name(table: TableInfo, column_name: str) -> ColumnInfo | None:
    for column in table.columns:
        if column.column_name == column_name:
            return column
    return None


def _score_inferred_relationship(
    source_column: ColumnInfo,
    target_table: TableInfo,
    target_column: ColumnInfo | None,
    coverage: float | None,
) -> tuple[float, dict[str, Any]]:
    score = 0.0
    evidence: dict[str, Any] = {}
    prefix = _reference_prefix(source_column.column_name)
    if prefix == target_table.table_name or target_table.table_name.endswith(prefix):
        score += 0.25
        evidence["name_match"] = True
    if target_column and source_column.data_type == target_column.data_type:
        score += 0.1
        evidence["type_match"] = True
    if source_column.column_name == f"{target_table.table_name}_id":
        score += 0.25
        evidence["table_prefix_match"] = True
    if target_table.primary_key and target_column and target_column.column_name == target_table.primary_key:
        score += 0.15
        evidence["target_primary_key"] = True
    if coverage is not None:
        evidence["data_coverage"] = round(float(coverage), 4)
        if coverage >= 0.95:
            score += 0.25
        elif coverage >= 0.7:
            score += 0.1
    return score, evidence


def _is_event_table(
    table: TableInfo,
    relationships: list[RelationshipCandidate],
) -> bool:
    if len(relationships) >= 2:
        return True
    table_name = table.table_name.lower()
    if any(hint in table_name for hint in _EVENT_HINTS) and relationships:
        return True
    column_names = {column.column_name.lower() for column in table.columns}
    return bool(relationships) and any(
        any(hint in column_name for hint in _EVENT_ATTR_HINTS)
        for column_name in column_names
    )


def _source_config_for_table(
    table: TableInfo,
    relationships: list[RelationshipCandidate],
    table_by_name: dict[str, TableInfo],
) -> dict[str, Any]:
    config: dict[str, Any] = {"primary_key": table.primary_key}
    if not relationships:
        config["table"] = table.table_name
        return config
    config["query"] = _joined_source_query(table, relationships, table_by_name)
    return config


def _joined_source_query(
    table: TableInfo,
    relationships: list[RelationshipCandidate],
    table_by_name: dict[str, TableInfo],
) -> str:
    base_alias = _table_alias(table.table_name)
    select_fields = [f"{base_alias}.*"]
    joins = []
    used_aliases = {base_alias}
    for relationship in relationships:
        target_table = table_by_name.get(relationship.target_table)
        if not target_table:
            continue
        display_field = _display_field(target_table)
        if not display_field:
            continue
        alias = _unique_alias(_table_alias(target_table.table_name), used_aliases)
        used_aliases.add(alias)
        output_name = _join_output_name(relationship, target_table, display_field)
        select_fields.append(
            f"{alias}.{display_field} AS {output_name}"
        )
        joins.append(
            "LEFT JOIN "
            f"{target_table.table_name} {alias} "
            f"ON {alias}.{relationship.target_column} = "
            f"{base_alias}.{relationship.source_column}"
        )
    select_sql = ",\n        ".join(select_fields)
    join_sql = "\n      ".join(joins)
    return (
        "SELECT\n"
        f"        {select_sql}\n"
        f"      FROM {table.table_name} {base_alias}"
        + (f"\n      {join_sql}" if join_sql else "")
    )


def _table_alias(table_name: str) -> str:
    parts = [part for part in table_name.split("_") if part]
    if len(parts) == 1:
        return parts[0][:1] or "t"
    return "".join(part[:1] for part in parts)


def _unique_alias(base_alias: str, used_aliases: set[str]) -> str:
    alias = base_alias
    index = 2
    while alias in used_aliases:
        alias = f"{base_alias}{index}"
        index += 1
    return alias


def _join_output_name(
    relationship: RelationshipCandidate,
    target_table: TableInfo,
    display_field: str,
) -> str:
    if display_field.startswith(f"{target_table.table_name}_"):
        return display_field
    prefix = _reference_prefix(relationship.source_column)
    return f"{prefix}_{display_field}"


def _entity_type_for_table(table: TableInfo, *, is_event: bool) -> str:
    if is_event:
        return _pascal_case(table.table_name)
    name = table.table_name.lower()
    if "company" in name or "enterprise" in name:
        return "Organization"
    if "project" in name:
        return "Project"
    if "person" in name or "user" in name:
        return "Person"
    return _pascal_case(table.table_name)


def _label_for_entity_type(entity_type: str) -> str:
    if entity_type == "Organization":
        return "organization"
    return re.sub(r"(?<!^)([A-Z])", r" \1", entity_type).lower()


def _entity_config_for_table(
    table: TableInfo,
    entity_type: str,
    relationships: list[RelationshipCandidate],
    table_by_name: dict[str, TableInfo],
    *,
    is_event: bool,
) -> dict[str, Any]:
    config: dict[str, Any] = {
        "source": table.table_name,
        "entity_type": entity_type,
        "id_field": table.primary_key,
        "metadata_fields": "*",
    }
    display_field = _display_field(table)
    if display_field and not is_event:
        config["name_field"] = display_field
        config["entity_name_template"] = "{" + display_field + "}"
    elif is_event:
        config["entity_name_template"] = _event_name_template(
            table,
            relationships,
            table_by_name,
        )
    else:
        config["entity_name_template"] = "{" + str(table.primary_key) + "}"
    config["description_template"] = _description_template(
        table,
        relationships,
        table_by_name,
        is_event=is_event,
    )
    return config


def _display_field(table: TableInfo) -> str | None:
    column_names = [column.column_name for column in table.columns]
    priorities = [
        f"{table.table_name}_name",
        "name",
        "title",
        "display_name",
    ]
    for field in priorities:
        if field in column_names:
            return field
    for column in table.columns:
        if "name" in column.column_name.lower() and _is_text_column(column):
            return column.column_name
    return None


def _is_text_column(column: ColumnInfo) -> bool:
    return column.data_type.lower() in _TEXT_TYPES or "char" in column.data_type.lower()


def _event_name_template(
    table: TableInfo,
    relationships: list[RelationshipCandidate],
    table_by_name: dict[str, TableInfo],
) -> str:
    display_placeholders = []
    for relationship in relationships:
        target_table = table_by_name.get(relationship.target_table)
        if not target_table:
            continue
        display_field = _display_field(target_table)
        if display_field:
            display_placeholders.append(
                "{" + _join_output_name(relationship, target_table, display_field) + "}"
            )
    if display_placeholders:
        return (
            f"{display_placeholders[0]} {table.table_name} "
            f"{display_placeholders[-1]} "
            + "({"
            + str(table.primary_key)
            + "})"
        )
    return "{" + str(table.primary_key) + "}"


def _description_template(
    table: TableInfo,
    relationships: list[RelationshipCandidate],
    table_by_name: dict[str, TableInfo],
    *,
    is_event: bool,
) -> str:
    fields = []
    if is_event:
        for relationship in relationships:
            target_table = table_by_name.get(relationship.target_table)
            if not target_table:
                continue
            display_field = _display_field(target_table)
            if display_field:
                output_name = _join_output_name(relationship, target_table, display_field)
                fields.append(f"{output_name}=" + "{" + output_name + "}")
    for column in table.columns:
        fields.append(f"{column.column_name}=" + "{" + column.column_name + "}")
    return "; ".join(fields)


def _relation_type(source_table: str, target_table: str) -> str:
    return f"{source_table}_{target_table}".upper()


def _relationship_as_metadata(relationship: RelationshipCandidate) -> dict[str, Any]:
    return {
        "source": relationship.source,
        "score": relationship.score,
        "decision": relationship.decision,
        "evidence": relationship.evidence,
    }


def _pascal_case(value: str) -> str:
    return "".join(part.capitalize() for part in value.split("_") if part)


def _load_postgres_tables(connection: Any, schema: str) -> list[TableInfo]:
    query = """
        SELECT
          c.table_name,
          c.column_name,
          c.data_type,
          c.is_nullable,
          obj_description((quote_ident(c.table_schema) || '.' || quote_ident(c.table_name))::regclass) AS table_comment,
          col_description((quote_ident(c.table_schema) || '.' || quote_ident(c.table_name))::regclass, c.ordinal_position) AS column_comment,
          pk.column_name AS primary_key
        FROM information_schema.columns c
        LEFT JOIN (
          SELECT ku.table_schema, ku.table_name, ku.column_name
          FROM information_schema.table_constraints tc
          JOIN information_schema.key_column_usage ku
            ON tc.constraint_name = ku.constraint_name
           AND tc.table_schema = ku.table_schema
           AND tc.table_name = ku.table_name
          WHERE tc.constraint_type = 'PRIMARY KEY'
        ) pk
          ON pk.table_schema = c.table_schema
         AND pk.table_name = c.table_name
        WHERE c.table_schema = %s
        ORDER BY c.table_name, c.ordinal_position
    """
    rows = connection.execute(query, (schema,)).fetchall()
    by_table: dict[str, dict[str, Any]] = {}
    for row in rows:
        (
            table_name,
            column_name,
            data_type,
            is_nullable,
            table_comment,
            column_comment,
            primary_key,
        ) = row
        table_data = by_table.setdefault(
            table_name,
            {
                "columns": [],
                "primary_key": primary_key,
                "comment": table_comment,
            },
        )
        if primary_key and not table_data.get("primary_key"):
            table_data["primary_key"] = primary_key
        table_data["columns"].append(
            ColumnInfo(
                table_name=table_name,
                column_name=column_name,
                data_type=data_type,
                is_nullable=(is_nullable == "YES"),
                comment=column_comment,
            )
        )
    return [
        TableInfo(
            table_name=table_name,
            columns=data["columns"],
            primary_key=data.get("primary_key"),
            comment=data.get("comment"),
        )
        for table_name, data in by_table.items()
    ]


async def _load_postgres_tables_asyncpg(connection: Any, schema: str) -> list[TableInfo]:
    query = """
        SELECT
          c.table_name,
          c.column_name,
          c.data_type,
          c.is_nullable,
          obj_description((quote_ident(c.table_schema) || '.' || quote_ident(c.table_name))::regclass) AS table_comment,
          col_description((quote_ident(c.table_schema) || '.' || quote_ident(c.table_name))::regclass, c.ordinal_position) AS column_comment,
          pk.column_name AS primary_key
        FROM information_schema.columns c
        LEFT JOIN (
          SELECT ku.table_schema, ku.table_name, ku.column_name
          FROM information_schema.table_constraints tc
          JOIN information_schema.key_column_usage ku
            ON tc.constraint_name = ku.constraint_name
           AND tc.table_schema = ku.table_schema
           AND tc.table_name = ku.table_name
          WHERE tc.constraint_type = 'PRIMARY KEY'
        ) pk
          ON pk.table_schema = c.table_schema
         AND pk.table_name = c.table_name
        WHERE c.table_schema = $1
        ORDER BY c.table_name, c.ordinal_position
    """
    rows = await connection.fetch(query, schema)
    by_table: dict[str, dict[str, Any]] = {}
    for row in rows:
        table_name = row["table_name"]
        column_name = row["column_name"]
        data_type = row["data_type"]
        is_nullable = row["is_nullable"]
        table_comment = row["table_comment"]
        column_comment = row["column_comment"]
        primary_key = row["primary_key"]
        table_data = by_table.setdefault(
            table_name,
            {
                "columns": [],
                "primary_key": primary_key,
                "comment": table_comment,
            },
        )
        if primary_key and not table_data.get("primary_key"):
            table_data["primary_key"] = primary_key
        table_data["columns"].append(
            ColumnInfo(
                table_name=table_name,
                column_name=column_name,
                data_type=data_type,
                is_nullable=(is_nullable == "YES"),
                comment=column_comment,
            )
        )
    return [
        TableInfo(
            table_name=table_name,
            columns=data["columns"],
            primary_key=data.get("primary_key"),
            comment=data.get("comment"),
        )
        for table_name, data in by_table.items()
    ]


def _load_postgres_foreign_keys(
    connection: Any,
    schema: str,
) -> list[RelationshipCandidate]:
    query = """
        SELECT
          kcu.table_name AS source_table,
          kcu.column_name AS source_column,
          ccu.table_name AS target_table,
          ccu.column_name AS target_column,
          tc.constraint_name
        FROM information_schema.table_constraints tc
        JOIN information_schema.key_column_usage kcu
          ON tc.constraint_name = kcu.constraint_name
         AND tc.table_schema = kcu.table_schema
        JOIN information_schema.constraint_column_usage ccu
          ON ccu.constraint_name = tc.constraint_name
         AND ccu.table_schema = tc.table_schema
        WHERE tc.constraint_type = 'FOREIGN KEY'
          AND tc.table_schema = %s
    """
    rows = connection.execute(query, (schema,)).fetchall()
    return [
        RelationshipCandidate(
            source_table=row[0],
            source_column=row[1],
            target_table=row[2],
            target_column=row[3],
            source="foreign_key",
            score=1.0,
            evidence={"constraint_name": row[4]},
            decision="auto_approved",
        )
        for row in rows
    ]


async def _load_postgres_foreign_keys_asyncpg(
    connection: Any,
    schema: str,
) -> list[RelationshipCandidate]:
    query = """
        SELECT
          kcu.table_name AS source_table,
          kcu.column_name AS source_column,
          ccu.table_name AS target_table,
          ccu.column_name AS target_column,
          tc.constraint_name
        FROM information_schema.table_constraints tc
        JOIN information_schema.key_column_usage kcu
          ON tc.constraint_name = kcu.constraint_name
         AND tc.table_schema = kcu.table_schema
        JOIN information_schema.constraint_column_usage ccu
          ON ccu.constraint_name = tc.constraint_name
         AND ccu.table_schema = tc.table_schema
        WHERE tc.constraint_type = 'FOREIGN KEY'
          AND tc.table_schema = $1
    """
    rows = await connection.fetch(query, schema)
    return [
        RelationshipCandidate(
            source_table=row["source_table"],
            source_column=row["source_column"],
            target_table=row["target_table"],
            target_column=row["target_column"],
            source="foreign_key",
            score=1.0,
            evidence={"constraint_name": row["constraint_name"]},
            decision="auto_approved",
        )
        for row in rows
    ]


def _load_relationship_coverage(
    connection: Any,
    schema: str,
    relationships: list[RelationshipCandidate],
    sample_limit: int,
) -> dict[tuple[str, str, str, str], float]:
    results: dict[tuple[str, str, str, str], float] = {}
    for relationship in relationships:
        query = coverage_query(
            schema,
            relationship.source_table,
            relationship.source_column,
            relationship.target_table,
            relationship.target_column,
            sample_limit,
        )
        total, matched = connection.execute(query).fetchone()
        key = (
            relationship.source_table,
            relationship.source_column,
            relationship.target_table,
            relationship.target_column,
        )
        results[key] = (float(matched) / float(total)) if total else 0.0
    return results


async def _load_relationship_coverage_asyncpg(
    connection: Any,
    schema: str,
    relationships: list[RelationshipCandidate],
    sample_limit: int,
) -> dict[tuple[str, str, str, str], float]:
    results: dict[tuple[str, str, str, str], float] = {}
    for relationship in relationships:
        query = coverage_query(
            schema,
            relationship.source_table,
            relationship.source_column,
            relationship.target_table,
            relationship.target_column,
            sample_limit,
        )
        row = await connection.fetchrow(query)
        total = row["total"] if row is not None else 0
        matched = row["matched"] if row is not None else 0
        key = (
            relationship.source_table,
            relationship.source_column,
            relationship.target_table,
            relationship.target_column,
        )
        results[key] = (float(matched) / float(total)) if total else 0.0
    return results
